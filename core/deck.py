"""PowerPoint (.pptx) report for one analysis run: simple, minimalist design, native (editable) charts.
Sections follow the app: brief, asset, tags, alignment, story (video), usage."""
import io
import math
import re
from datetime import date

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_TICK_LABEL_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

from core.scoring import TagScore
from core.video import fmt_time

APP_NAME = "Creative Content Analyzer"
GITHUB_URL = "https://github.com/cesarbrea/content-analyzer"
CONTACT_EMAIL = "cesar.brea@alixpartners.com"

INK = RGBColor(0x1F, 0x29, 0x37)
MUTED = RGBColor(0x6B, 0x72, 0x80)
ACCENT = RGBColor(0x1F, 0x3A, 0x5F)
GOOD = RGBColor(0x2E, 0x7D, 0x5B)
BAD = RGBColor(0xC0, 0x50, 0x4D)
FONT = "Calibri"

SLIDE_W, SLIDE_H = Inches(13.333), Inches(7.5)
MARGIN = 0.7                      # inches
BODY_TOP, BODY_BOTTOM = 1.45, 6.75
BODY_W = 13.333 - 2 * MARGIN
TAGS_ONE_COLUMN_MAX = 12          # more tags than this are split into two charts side by side
BEATS_PER_SLIDE = 6               # 2 columns x 3 rows


def footer_text(org: str) -> str:
    """Footer line shared by the PDF and the deck; the organization part is left out when none was submitted."""
    parts = ([f"Creative Content Analysis for {org.strip()}"] if org.strip() else []) + [GITHUB_URL, CONTACT_EMAIL]
    return "   |   ".join(parts)


def _t(text: str) -> str:
    """python-pptx rejects control characters that can appear in model output."""
    return "".join(ch for ch in text if ch == "\n" or ch >= " ")


# ---------------------------------------------------------------- text fitting (PowerPoint doesn't auto-shrink)

def _lines(paras: list[str], width_in: float, size: int) -> int:
    chars_per_line = max(10, int(width_in * 72 / (size * 0.5)))
    return sum(max(1, math.ceil(len(p) / chars_per_line)) for p in paras)


def fit_size(paras: list[str], width_in: float, height_in: float, sizes=(18, 16, 14, 13, 12, 11, 10),
             gap_pt: float = 6) -> int | None:
    """Largest font size at which the paragraphs (estimated) fit in the box, or None."""
    for size in sizes:
        need = _lines(paras, width_in, size) * size * 1.2 + len(paras) * gap_pt
        if need <= height_in * 72:
            return size
    return None


def page_text(paras: list[str], width_in: float, height_in: float, min_size: int = 12) -> list[list[str]]:
    """Split paragraphs over as many slides as needed so each page fits at >= min_size."""
    pages, cur = [], []
    for p in paras:
        # a single paragraph too long for a page is split on sentence-ish boundaries
        chunks = [p]
        while chunks and fit_size([chunks[-1]], width_in, height_in, sizes=(min_size,)) is None:
            long = chunks.pop()
            cut = long.rfind(". ", 0, len(long) // 2 + 200)
            cut = cut + 2 if cut > 0 else len(long) // 2
            chunks += [long[:cut].strip(), long[cut:].strip()]
        for c in chunks:
            if cur and fit_size(cur + [c], width_in, height_in, sizes=(min_size,)) is None:
                pages.append(cur)
                cur = []
            cur.append(c)
    if cur:
        pages.append(cur)
    return pages or [[]]


# ---------------------------------------------------------------- slide building blocks

class Deck:
    def __init__(self, org: str = ""):
        self.org = org.strip()
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = SLIDE_W, SLIDE_H
        self.sections: list[tuple[str, int]] = []   # (name, first slide number)

    # -- slides
    def slide(self, title: str, section: str | None = None):
        s = self.prs.slides.add_slide(self.prs.slide_layouts[5])  # "Title Only": a real title placeholder
        n = len(self.prs.slides)
        if section and not any(name == section for name, _ in self.sections):
            self.sections.append((section, n))
        t = s.shapes.title
        t.left, t.top, t.width, t.height = Inches(MARGIN), Inches(0.45), Inches(BODY_W), Inches(0.8)
        t.text_frame.margin_left = 0
        t.text_frame.text = title
        t.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = t.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.LEFT
        self._font(p.runs[0], 30, True, INK)
        rule = s.shapes.add_shape(1, Inches(MARGIN), Inches(1.28), Inches(1.2), Emu(38100))  # short accent rule
        rule.fill.solid()
        rule.fill.fore_color.rgb = ACCENT
        rule.line.fill.background()
        self._footer(s)
        return s

    def _footer(self, s):
        box = s.shapes.add_textbox(Inches(MARGIN), Inches(6.95), Inches(10.6), Inches(0.35))
        box.text_frame.margin_left = 0
        p = box.text_frame.paragraphs[0]
        r = p.add_run()
        r.text = _t(footer_text(self.org))
        self._font(r, 10, False, MUTED)
        num = s.shapes.add_textbox(Inches(13.333 - MARGIN - 1), Inches(6.95), Inches(1), Inches(0.35))
        num.text_frame.margin_right = 0
        p = num.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.RIGHT
        fld = p._p.makeelement(qn("a:fld"), {"id": "{B6F15528-21DE-4FAA-801E-634DDDAF4B2B}", "type": "slidenum"})
        rpr = fld.makeelement(qn("a:rPr"), {"lang": "en-US", "sz": "1000"})
        fill = rpr.makeelement(qn("a:solidFill"), {})
        fill.append(fill.makeelement(qn("a:srgbClr"), {"val": "6B7280"}))
        rpr.append(fill)
        fld.append(rpr)
        t = fld.makeelement(qn("a:t"), {})
        t.text = str(len(self.prs.slides))
        fld.append(t)
        p._p.append(fld)

    @staticmethod
    def _font(run, size, bold=False, color=INK, italic=False):
        run.font.name = FONT
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.italic = italic
        run.font.color.rgb = color

    # -- text
    def text(self, s, x, y, w, h, paras, size=14, color=INK, gap=6, bullet=False):
        """paras: list of str, or of [(text, bold), ...] run lists."""
        box = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = 0
        for i, para in enumerate(paras):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.space_after = Pt(gap)
            runs = [(para, False)] if isinstance(para, str) else para
            if bullet:
                runs = [("•  ", False)] + list(runs)
            for text, bold in runs:
                r = p.add_run()
                r.text = _t(text)
                self._font(r, size, bold, color)
        return box

    def picture(self, s, jpeg: bytes, x, y, max_w, max_h):
        from PIL import Image
        w, h = Image.open(io.BytesIO(jpeg)).size
        k = min(max_w / w, max_h / h)
        pw, ph = w * k, h * k
        s.shapes.add_picture(io.BytesIO(jpeg), Inches(x), Inches(y), Inches(pw), Inches(ph))
        return pw, ph

    # -- charts
    def bar_chart(self, s, x, y, w, h, cats, vals, fmt, vmin, vmax, colors=None, label_size=12):
        data = CategoryChartData()
        data.categories = cats
        data.add_series("Value", vals)
        gf = s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(x), Inches(y), Inches(w), Inches(h), data)
        ch = gf.chart
        ch.has_legend = False
        ch.has_title = False
        ch.font.name = FONT
        ch.font.size = Pt(label_size)
        ch.font.color.rgb = INK
        va, ca = ch.value_axis, ch.category_axis
        va.minimum_scale, va.maximum_scale = vmin, vmax
        va.visible = False
        va.has_major_gridlines = False
        ca.tick_label_position = XL_TICK_LABEL_POSITION.LOW   # labels stay left even for negative bars
        ca.format.line.color.rgb = MUTED
        ca.has_major_gridlines = False
        ca.reverse_order = True                                # first item at the top
        plot = ch.plots[0]
        plot.gap_width = 45
        plot.has_data_labels = True
        dl = plot.data_labels
        dl.number_format, dl.number_format_is_linked = fmt, False
        dl.position = XL_LABEL_POSITION.OUTSIDE_END
        dl.font.size = Pt(label_size)
        ser = plot.series[0]
        ser.invert_if_negative = False
        ser.format.fill.solid()
        ser.format.fill.fore_color.rgb = ACCENT
        if colors:
            for i, c in enumerate(colors):
                pt = ser.points[i]
                pt.format.fill.solid()
                pt.format.fill.fore_color.rgb = c
                dpt = ser._element.findall(qn("c:dPt"))[-1]
                inv = dpt.makeelement(qn("c:invertIfNegative"), {"val": "0"})
                dpt.find(qn("c:idx")).addnext(inv)   # without this, negative bars render unfilled
        return gf


# ---------------------------------------------------------------- the report

def _nearest(frames: list[tuple[float, bytes]], t: float) -> bytes | None:
    return min(frames, key=lambda f: abs(f[0] - t))[1] if frames else None


def build_deck(result: dict, shown: list[TagScore], product_desc: str, brief: str, org: str = "") -> bytes:
    """`shown` is the tag list currently on screen; `brief` is used when the result has no alignment."""
    D = Deck(org)
    alignment = result.get("alignment")
    brief_used = (alignment["brief"] if alignment else brief).strip()
    story = result.get("story")
    story_frames = result.get("story_frames") or []

    # cover
    cover = D.prs.slides.add_slide(D.prs.slide_layouts[6])
    D.text(cover, 1.0, 2.3, 11.3, 1.4, [[(f"{APP_NAME} Report", True)]], size=44, color=INK)
    y = 3.9
    if D.org:
        D.text(cover, 1.0, y, 11.3, 0.6, [[(D.org, False)]], size=26, color=ACCENT)
        y += 0.8
    D.text(cover, 1.0, y, 11.3, 1.0, [date.today().strftime("%B %d, %Y").replace(" 0", " ")], size=16, color=MUTED)
    bar = cover.shapes.add_shape(1, Inches(1.0), Inches(3.75), Inches(1.6), Emu(50800))
    bar.fill.solid()
    bar.fill.fore_color.rgb = ACCENT
    bar.line.fill.background()

    contents = D.slide("Contents")

    # product description
    desc_paras = [p for p in product_desc.strip().splitlines() if p.strip()]
    for i, page in enumerate(page_text(desc_paras, BODY_W, BODY_BOTTOM - BODY_TOP) if desc_paras else []):
        s = D.slide("Product description" + (" (continued)" if i else ""), "Product description")
        D.text(s, MARGIN, BODY_TOP, BODY_W, BODY_BOTTOM - BODY_TOP, page,
               size=fit_size(page, BODY_W, BODY_BOTTOM - BODY_TOP) or 12)

    # creative brief: one slide, a rounded box per element with its label above and a bullet per sentence
    if brief_used:
        _brief_slide(D, brief_used)

    # asset evaluated
    s = D.slide("Asset evaluated", "Asset evaluated")
    pw, ph = D.picture(s, result["image_jpeg"], MARGIN, BODY_TOP, 6.0, 5.0)
    tx = MARGIN + 6.0 + 0.5
    tw = 13.333 - MARGIN - tx
    meta = result["kind"].capitalize()
    m = result.get("video_meta") or {}
    if m.get("duration"):
        meta += f" · {m['duration']:.1f}s · {m['width']}×{m['height']}"
    D.text(s, tx, BODY_TOP, tw, 0.4, [[("Filename", True)]], size=14, color=MUTED)
    D.text(s, tx, BODY_TOP + 0.4, tw, 0.8, [result["filename"], meta], size=16, gap=2)
    D.text(s, tx, BODY_TOP + 1.5, tw, 0.4, [[("Caption", True)]], size=14, color=MUTED)
    cap = [result["caption"]]
    D.text(s, tx, BODY_TOP + 1.9, tw, 3.3, cap, size=fit_size(cap, tw, 3.3) or 12)

    # tags detected
    if shown:
        s = D.slide("Tags detected", "Tags detected")
        D.text(s, MARGIN, 1.35, BODY_W, 0.35,
               [f"Certainty for the top {len(shown)} tags (agreement across {result['k']} runs, weighted with "
                "the model's self-confidence)"], size=12, color=MUTED)
        cols = [shown] if len(shown) <= TAGS_ONE_COLUMN_MAX else [shown[:math.ceil(len(shown) / 2)],
                                                                  shown[math.ceil(len(shown) / 2):]]
        cw = (BODY_W - 0.4) / len(cols)
        for i, col in enumerate(cols):
            D.bar_chart(s, MARGIN + i * (cw + 0.4), 1.8, cw, 5.0, [t.tag for t in col],
                        [t.certainty for t in col], "0%", 0, 1.15)

    # alignment
    if alignment:
        _summary_slide(D, "Creative brief alignment summary", alignment["overall"],
                       "Overall alignment: the mean of the four dimensions.",
                       [n.capitalize() for n in alignment["dimensions"]],
                       [v["score"] for v in alignment["dimensions"].values()])

        # evaluation and suggestions
        lw, gap = 6.3, 0.5
        rw = BODY_W - lw - gap
        h = BODY_BOTTOM - BODY_TOP - 0.45
        evals = [[(f"{n.capitalize()} ({v['score']:+.2f}): ", True), (v["rationale"], False)]
                 for n, v in alignment["dimensions"].items()]
        eval_plain = [f"{n.capitalize()} ({v['score']:+.2f}): {v['rationale']}" for n, v in alignment["dimensions"].items()]
        sugg = list(alignment["suggestions"])
        _two_column_slide(D, "Creative brief alignment: evaluation and suggestions",
                                    "Creative brief alignment: evaluation and suggestions", "Evaluation", "Suggestions",
                                    evals, eval_plain, sugg, lw, rw, gap, h)

    # story (video)
    if story:
        if story.get("scores"):
            sc = story["scores"]
            vals = [v["score"] for v in sc.values()]
            _summary_slide(D, "Story alignment summary", sum(vals) / len(vals),
                           "Overall story alignment: the mean of the four story dimensions. Shown separately from the "
                           "overall brief alignment.",
                           [d.replace("_", " ").capitalize() for d in sc], vals)
        beats = story["beats"]
        for i in range(0, max(1, len(beats)), BEATS_PER_SLIDE):
            s = D.slide("Story detected" + (" (continued)" if i else ""), "Story detected")
            top = BODY_TOP
            if i == 0:
                D.text(s, MARGIN, 1.4, BODY_W, 0.7, [story["logline"]], size=fit_size([story["logline"]], BODY_W, 0.7) or 12,
                       color=MUTED)
                top = 2.2
            chunk = beats[i:i + BEATS_PER_SLIDE]
            _thumb_grid(D, s, top, [
                (_nearest(story_frames, (b["start"] + b["end"]) / 2), f"{fmt_time(b['start'])}–{fmt_time(b['end'])}",
                 b["summary"]) for b in chunk])
        chars = story["characters"]
        for i in range(0, len(chars), BEATS_PER_SLIDE):
            s = D.slide("Story detected: characters" + (" (continued)" if i else ""), "Story detected")
            _thumb_grid(D, s, BODY_TOP, [
                (_nearest(story_frames, c["first_seen"]), c["label"],
                 f"{c['role']}; {c['presence']}{', speaks' if c['speaks'] else ''}. {c['description']}")
                for c in chars[i:i + BEATS_PER_SLIDE]])
        s = D.slide("Story detected: arc and limits", "Story detected")
        paras = [[("Arc: ", True), (story["arc"], False)]]
        plain = [f"Arc: {story['arc']}"]
        if story["gaps"]:
            paras.append([("What could not be determined", True)])
            plain.append("What could not be determined")
            for g in story["gaps"]:
                paras.append([("•  " + g, False)])
                plain.append("•  " + g)
        paras.append([("Inferred from the sampled frames and the transcript; music, sound effects and action "
                       "between frames are not seen.", False)])
        plain.append("Inferred from the sampled frames and the transcript; music, sound effects and action "
                     "between frames are not seen.")
        D.text(s, MARGIN, BODY_TOP, BODY_W, BODY_BOTTOM - BODY_TOP, paras,
               size=fit_size(plain, BODY_W, BODY_BOTTOM - BODY_TOP) or 11)

        if story.get("scores"):
            lw, gap = 6.3, 0.5
            rw = BODY_W - lw - gap
            h = BODY_BOTTOM - BODY_TOP - 0.45
            evals = [[(f"{d.replace('_', ' ').capitalize()} ({v['score']:+.2f}): ", True), (v["rationale"], False)]
                     for d, v in story["scores"].items()]
            plain = [f"{d.replace('_', ' ').capitalize()} ({v['score']:+.2f}): {v['rationale']}"
                     for d, v in story["scores"].items()]
            _two_column_slide(D, "Story evaluation and suggestions", "Story evaluation and suggestions", "Evaluation", "Suggestions",
                              evals, plain, list(story["suggestions"]), lw, rw, gap, h)

    # usage
    s = D.slide("Usage", "Usage")
    u = result["usage"]
    rows = [("Model", result["model"]), ("Runs per asset (K)", str(result["k"])),
            ("Tokens in / out", f"{u['input_tokens']:,} / {u['output_tokens']:,}"),
            ("Cost of this session", f"${result['cost_usd']:.4f}")]
    for i, (label, value) in enumerate(rows):
        y = BODY_TOP + 0.2 + i * 0.75
        D.text(s, MARGIN, y, 3.5, 0.5, [[(label, True)]], size=18, color=MUTED)
        D.text(s, MARGIN + 3.6, y, 7, 0.5, [value], size=20)

    # contents, filled last so it can show page numbers
    names, nums = [n for n, _ in D.sections], [str(p) for _, p in D.sections]
    D.text(contents, MARGIN, BODY_TOP, 8, 5, [f"{i + 1}.  {n}" for i, n in enumerate(names)], size=20, gap=14)
    num_box = D.text(contents, MARGIN + 8, BODY_TOP, 1.5, 5, nums, size=20, color=MUTED, gap=14)
    for p in num_box.text_frame.paragraphs:
        p.alignment = PP_ALIGN.RIGHT

    buf = io.BytesIO()
    D.prs.save(buf)
    return buf.getvalue()


_ABBREV = ("e.g.", "i.e.", "vs.", "etc.", "approx.", "inc.", "no.", "st.", "mr.", "mrs.", "dr.")
_LABEL = re.compile(r"^\s*(?:[#*\-\u2022]+\s*)?\**([A-Za-z][^:\n]{1,40}?)\**\s*:\s*(.*)$")


def split_sentences(text: str) -> list[str]:
    """Split on sentence ends, but not after common abbreviations (e.g., i.e., vs.)."""
    parts, start = [], 0
    for m in re.finditer(r"[.!?][\"')\u201d]?\s+(?=[A-Z0-9\"(\u201c])", text):
        chunk = text[start:m.end()].strip()
        if chunk.lower().endswith(_ABBREV):
            continue
        parts.append(chunk)
        start = m.end()
    tail = text[start:].strip()
    if tail:
        parts.append(tail)
    return parts


def parse_brief(text: str) -> list[tuple[str, list[str]]]:
    """Split a brief into (label, sentences). Two styles are understood:
    - headed: "Label:" alone on a line, text on the following lines (what the app's brief drafting produces);
      only those lines are labels, so a colon inside a sentence ("Show moments: a customer...") is left alone;
    - inline: "Label: text" on one line.
    A brief with no recognisable labels becomes one section called "Creative brief"."""
    lines = [ln for ln in text.splitlines() if ln.strip()]
    matches = [_LABEL.match(ln) for ln in lines]
    headed = sum(1 for m in matches if m and not m.group(2).strip() and len(m.group(1).split()) <= 6) >= 2
    sections: list[tuple[str, list[str]]] = []
    for line, m in zip(lines, matches):
        is_label = bool(m) and len(m.group(1).split()) <= 6 and (not m.group(2).strip() or not headed)
        if is_label:
            sections.append((m.group(1).strip(), []))
            line = m.group(2)
        elif not sections:
            sections.append(("Creative brief", []))
        if line.strip():
            sections[-1][1].extend(split_sentences(line.strip()))
    sections = [(label, sents) for label, sents in sections if sents]
    if len(sections) < 2:   # no real structure: show everything as one element
        return [("Creative brief", split_sentences(" ".join(text.split())))] if text.strip() else []
    return sections


def _bullets(tf, sentences, size, color=INK):
    """One paragraph per sentence, with a real bullet and a hanging indent."""
    for i, sent in enumerate(sentences):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT   # shapes centre text by default
        p.space_after = Pt(4)
        r = p.add_run()
        r.text = _t(sent)
        Deck._font(r, size, False, color)
        ppr = p._p.get_or_add_pPr()
        ppr.set("marL", str(Inches(0.18)))
        ppr.set("indent", str(-Inches(0.18)))
        ppr.append(ppr.makeelement(qn("a:buChar"), {"char": "\u2022"}))


def _brief_slide(D, brief: str):
    sections = parse_brief(brief)
    n = len(sections)
    rows = 1 if n <= 3 else 2 if n <= 8 else 3
    cols = math.ceil(n / rows)
    gap_x, gap_y, label_h = 0.25, 0.2, 0.34
    cell_w = (BODY_W - (cols - 1) * gap_x) / cols
    cell_h = (BODY_BOTTOM - BODY_TOP - (rows - 1) * gap_y) / rows
    box_h = cell_h - label_h
    inner_w = cell_w - 0.24 - 0.18   # box margins and the bullet indent
    inner_h = box_h - 0.2
    # one font size for every box: the largest at which the fullest box fits
    size = min(fit_size(sents, inner_w, inner_h, sizes=(16, 15, 14, 13, 12, 11, 10, 9, 8), gap_pt=4) or 8
               for _, sents in sections)
    s = D.slide("Creative brief", "Creative brief")
    for i, (label, sents) in enumerate(sections):
        x = MARGIN + (i % cols) * (cell_w + gap_x)
        y = BODY_TOP + (i // cols) * (cell_h + gap_y)
        D.text(s, x, y, cell_w, label_h, [[(label, True)]], size=13, color=ACCENT, gap=0)
        box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y + label_h), Inches(cell_w), Inches(box_h))
        box.adjustments[0] = 0.06
        box.fill.solid()
        box.fill.fore_color.rgb = RGBColor(0xF3, 0xF4, 0xF6)
        box.line.color.rgb = RGBColor(0xD1, 0xD5, 0xDB)
        box.line.width = Pt(1)
        box.shadow.inherit = False
        tf = box.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.TOP
        tf.margin_left = tf.margin_right = Inches(0.12)
        tf.margin_top = tf.margin_bottom = Inches(0.1)
        _bullets(tf, sents, size)
    return s


def _summary_slide(D, title, overall, blurb, names, scores):
    """Big overall number on the left, a green/red bar per dimension on the right."""
    s = D.slide(title, title)
    D.text(s, MARGIN, BODY_TOP, 3.4, 1.2, [[(f"{overall:+.2f}", True)]], size=54, color=ACCENT)
    D.text(s, MARGIN, BODY_TOP + 1.2, 3.4, 2.8,
           [blurb, "-1 contradicts the brief, 0 is absent or not addressed, +1 fully conforms."], size=13,
           color=MUTED, gap=8)
    D.bar_chart(s, 4.4, BODY_TOP, 13.333 - MARGIN - 4.4, 4.9, names, scores, "+0.00;-0.00;0.00", -1.2, 1.2,
                colors=[GOOD if v >= 0 else BAD for v in scores], label_size=14)
    return s


def _two_column_slide(D, title, section, left_head, right_head, left, left_plain, right, lw, rw, gap, h):
    """Evaluation on the left, suggestions on the right; splits to two slides if it can't fit at >= 11 pt."""
    size = min(fit_size(left_plain, lw, h, sizes=(16, 15, 14, 13, 12, 11)) or 0,
               fit_size([f"•  {r}" for r in right], rw, h, sizes=(16, 15, 14, 13, 12, 11)) or 0)
    if size:
        s = D.slide(title, section)
        D.text(s, MARGIN, BODY_TOP, lw, 0.4, [[(left_head, True)]], size=14, color=MUTED)
        D.text(s, MARGIN, BODY_TOP + 0.45, lw, h, left, size=size, gap=8)
        D.text(s, MARGIN + lw + gap, BODY_TOP, rw, 0.4, [[(right_head, True)]], size=14, color=MUTED)
        D.text(s, MARGIN + lw + gap, BODY_TOP + 0.45, rw, h, right, size=size, gap=8, bullet=True)
        return s
    s = D.slide(title, section)
    D.text(s, MARGIN, BODY_TOP, BODY_W, BODY_BOTTOM - BODY_TOP, left, size=fit_size(left_plain, BODY_W, 5.2) or 11, gap=8)
    s = D.slide(title + " (continued)", section)
    D.text(s, MARGIN, BODY_TOP, BODY_W, BODY_BOTTOM - BODY_TOP, right,
           size=fit_size(right, BODY_W, 5.2) or 11, gap=8, bullet=True)
    return s


def _thumb_grid(D, s, top, items):
    """items: [(jpeg or None, heading, text)]; two columns, three rows."""
    colw = (BODY_W - 0.5) / 2
    rowh = (BODY_BOTTOM - top) / 3
    thumb_w, thumb_h = 1.9, rowh - 0.25
    for i, (jpeg, head, text) in enumerate(items):
        col, row = i % 2, i // 2
        x = MARGIN + col * (colw + 0.5)
        y = top + row * rowh
        pw = 0
        if jpeg:
            pw, _ = D.picture(s, jpeg, x, y, thumb_w, thumb_h)
        tx = x + (thumb_w if jpeg else 0) + 0.2
        tw = colw - (thumb_w if jpeg else 0) - 0.2
        plain = [head, text]
        size = fit_size(plain, tw, thumb_h, sizes=(13, 12, 11, 10)) or 10
        D.text(s, tx, y - 0.05, tw, thumb_h + 0.1, [[(head, True)], text], size=size, gap=2)
