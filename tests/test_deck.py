import io

from PIL import Image
from pptx import Presentation

from core.deck import build_deck, fit_size, page_text
from core.scoring import TagScore


def _jpg(color=(40, 80, 140), size=(320, 180)):
    b = io.BytesIO()
    Image.new("RGB", size, color).save(b, "JPEG")
    return b.getvalue()


def _result(video: bool, with_brief: bool):
    dims = {d: {"score": s, "rationale": "Because."} for d, s in zip(["customer", "context", "product", "style"],
                                                                   [0.5, -0.5, 0.0, 1.0])}
    r = {"filename": "ad.mp4" if video else "pic.png", "kind": "video" if video else "image", "caption": "A caption.",
         "alignment": {"brief": "Objective: sell things.\nAudience: everyone.", "overall": 0.25, "dimensions": dims,
                       "suggestions": ["Do this", "Do that"]} if with_brief else None,
         "story": None, "story_frames": [], "image_jpeg": _jpg(), "video_meta": {},
         "cost_usd": 0.05, "k": 3, "model": "claude-sonnet-5", "usage": {"input_tokens": 100, "output_tokens": 50}}
    if video:
        sc = {d: {"score": 0.2, "rationale": "Why."} for d in ["plot", "characters", "structure", "product_timing"]}
        r["story"] = {"logline": "Things happen.", "arc": "Start, middle, end.", "gaps": ["music"],
                      "beats": [{"start": i, "end": i + 1, "summary": f"beat {i}", "evidence": ""} for i in range(8)],
                      "characters": [{"label": "host", "role": "lead", "description": "a person", "first_seen": 0.0,
                                      "presence": "central", "speaks": True}],
                      "scores": sc if with_brief else None, "suggestions": ["Show product early"] if with_brief else None}
        r["story_frames"] = [(i * 0.75, _jpg()) for i in range(12)]
        r["video_meta"] = {"duration": 8.0, "width": 1920, "height": 1080}
    return r


def _tags(n):
    return [TagScore(f"tag {i}", "visual", 1 - i / 40, 0.8, 1 - i / 40, 3, "x") for i in range(n)]


def _titles(pptx_bytes):
    prs = Presentation(io.BytesIO(pptx_bytes))
    return [s.shapes.title.text_frame.text if s.shapes.title is not None else "(cover)" for s in prs.slides], prs


def test_video_with_brief_has_every_section_in_order():
    titles, prs = _titles(build_deck(_result(True, True), _tags(20), "A product.", ""))
    assert titles[0] == "(cover)" and titles[1] == "Contents"
    wanted = ["Product description", "Creative brief", "Asset evaluated", "Tags detected", "Creative brief alignment summary",
              "Creative brief alignment: evaluation and suggestions", "Story alignment summary", "Story detected",
              "Story evaluation and suggestions", "Usage"]
    positions = [titles.index(w) for w in wanted]
    assert positions == sorted(positions) and titles[-1] == "Usage"
    cover = " ".join(sh.text_frame.text for sh in prs.slides[0].shapes if sh.has_text_frame)
    assert "Creative Content Analyzer Report" in cover and "A product." not in cover
    descr = [sl for sl in prs.slides if sl.shapes.title is not None and sl.shapes.title.text_frame.text == "Product description"][0]
    assert "A product." in _all_text(descr) and titles.index("Product description") == 2


def test_image_without_brief_skips_brief_alignment_and_story():
    titles, _ = _titles(build_deck(_result(False, False), _tags(5), "", ""))
    assert titles == ["(cover)", "Contents", "Asset evaluated", "Tags detected", "Usage"]


def test_many_tags_split_into_two_charts_and_few_stay_in_one():
    def charts(n):
        prs = Presentation(io.BytesIO(build_deck(_result(False, False), _tags(n), "", "")))
        tags_slide = [s for s in prs.slides if s.shapes.title is not None and s.shapes.title.text_frame.text == "Tags detected"][0]
        return sum(1 for sh in tags_slide.shapes if sh.has_chart)
    assert charts(10) == 1 and charts(20) == 2


def test_story_without_brief_has_no_story_evaluation():
    titles, _ = _titles(build_deck(_result(True, False), _tags(5), "", ""))
    assert "Story detected" in titles and "Story evaluation and suggestions" not in titles
    assert "Creative brief" not in titles


def _brief_boxes(slide):
    return [sh for sh in slide.shapes if sh.shape_type == 1 and sh.has_text_frame and sh.text_frame.text]


def test_eight_element_brief_is_two_slides_top_row_then_bottom_row():
    r = _result(False, True)
    r["alignment"]["brief"] = "\n".join(f"Section {i}:\n" + "A sentence about the audience and tone of the campaign. " * 6
                                        for i in range(8))
    titles, prs = _titles(build_deck(r, _tags(5), "", ""))
    assert titles.count("Creative brief") == 1 and titles.count("Creative brief (continued)") == 1
    first, second = prs.slides[titles.index("Creative brief")], prs.slides[titles.index("Creative brief (continued)")]
    assert titles.index("Creative brief (continued)") == titles.index("Creative brief") + 1
    b1, b2 = _brief_boxes(first), _brief_boxes(second)
    assert [b.text_frame.text[:2] for b in b1] == ["A "] * 4 and len(b1) == 4 and len(b2) == 4
    # the first slide holds elements 0-3 (one row, left to right) and the second 4-7, at the same height
    labels1 = [sh.text_frame.text for sh in first.shapes if sh.has_text_frame and sh.text_frame.text.startswith("Section")]
    labels2 = [sh.text_frame.text for sh in second.shapes if sh.has_text_frame and sh.text_frame.text.startswith("Section")]
    assert labels1 == [f"Section {i}" for i in range(4)] and labels2 == [f"Section {i}" for i in range(4, 8)]
    assert len({b.top for b in b1}) == 1 and {b.top for b in b1} == {b.top for b in b2}
    assert [b.left for b in b1] == sorted(b.left for b in b1)
    # a bullet per sentence, with real bullet formatting, and the label above (not inside) each box
    assert len(b1[0].text_frame.paragraphs) == 6
    assert all(p._p.pPr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}buChar") is not None
               for p in b1[0].text_frame.paragraphs)
    lab = [sh for sh in first.shapes if sh.has_text_frame and sh.text_frame.text.startswith("Section")
           and sh.shape_type != 1]
    assert len(lab) == 4 and all(l.top < b.top for l, b in zip(lab, b1))


def test_small_brief_stays_on_one_slide():
    r = _result(False, True)
    r["alignment"]["brief"] = "Aim:\nOne thing.\n\nBuyer:\nAnother thing.\n\nChannel:\nA third."
    titles, prs = _titles(build_deck(r, _tags(5), "", ""))
    assert titles.count("Creative brief") == 1 and "Creative brief (continued)" not in titles
    assert len(_brief_boxes(prs.slides[titles.index("Creative brief")])) == 3


def test_eleven_element_brief_is_two_slides_with_rows_split_evenly():
    r = _result(False, True)
    r["alignment"]["brief"] = "\n".join(f"Section {i}:\nA short sentence." for i in range(11))
    titles, prs = _titles(build_deck(r, _tags(5), "", ""))
    n1 = len(_brief_boxes(prs.slides[titles.index("Creative brief")]))
    n2 = len(_brief_boxes(prs.slides[titles.index("Creative brief (continued)")]))
    assert n1 + n2 == 11 and n2 > 0


def test_parse_brief_sections_and_sentences():
    from core.deck import parse_brief, split_sentences
    got = parse_brief("Objective:\nDrive trial. Grow repeat visits.\n\nTarget audience: Owners, e.g. busy parents. They value time.")
    assert got == [("Objective", ["Drive trial.", "Grow repeat visits."]),
                   ("Target audience", ["Owners, e.g. busy parents.", "They value time."])]
    assert parse_brief("No labels here. Just prose.") == [("Creative brief", ["No labels here.", "Just prose."])]
    assert parse_brief("  ") == []
    assert split_sentences("Costs $5.50 per unit. Cheap.") == ["Costs $5.50 per unit.", "Cheap."]


def test_fit_and_paging_helpers():
    assert fit_size(["short"], 10, 5) == 18
    assert fit_size(["x" * 20000], 10, 5) is None
    pages = page_text(["word " * 400] * 6, 11.9, 5.3)
    assert len(pages) > 1 and sum(len(p) for p in pages) >= 6


def _all_text(slide):
    return " ".join(sh.text_frame.text for sh in slide.shapes if sh.has_text_frame)


def test_organization_is_cover_subtitle_and_in_every_footer():
    _, prs = _titles(build_deck(_result(True, True), _tags(5), "A product.", "", org="Acme Marine"))
    assert "Acme Marine" in _all_text(prs.slides[0])
    for slide in list(prs.slides)[1:]:
        assert "Creative Content Analysis for Acme Marine" in _all_text(slide)
        assert "cesar.brea@alixpartners.com" in _all_text(slide)


def test_no_organization_leaves_it_out():
    _, prs = _titles(build_deck(_result(False, False), _tags(5), "", "", org="  "))
    assert all("Creative Content Analysis for" not in _all_text(s) for s in prs.slides)


def test_story_summary_comes_right_before_story_detected_and_shows_the_mean():
    titles, prs = _titles(build_deck(_result(True, True), _tags(5), "", ""))
    i = titles.index("Story alignment summary")
    assert titles[i + 1] == "Story detected"
    assert "+0.20" in _all_text(prs.slides[i])   # mean of four 0.2 scores


def test_pdf_has_organization_subtitle_and_footer():
    import io as _io
    from pypdf import PdfReader
    from core.report import build_report
    r = _result(False, True)
    pdf = build_report(r, _tags(5), "A product.", "", org="Acme Marine")
    reader = PdfReader(_io.BytesIO(pdf))
    text = reader.pages[0].extract_text()
    assert "Acme Marine" in text and "Creative Content Analysis for Acme Marine" in text
    assert "Creative Content Analysis for" not in PdfReader(_io.BytesIO(build_report(r, _tags(5), "", ""))).pages[0].extract_text()


def test_no_product_description_means_no_description_slide():
    titles, _ = _titles(build_deck(_result(False, True), _tags(5), "  ", ""))
    assert "Product description" not in titles


def test_colon_inside_a_sentence_is_not_a_label_in_a_headed_brief():
    from core.deck import parse_brief
    got = parse_brief("Setting:" + chr(10) + "Show real-world moments: a customer pulling in. Use bright visuals." + chr(10) + chr(10) + "Tone:" + chr(10) + "Friendly.")
    assert [label for label, _ in got] == ["Setting", "Tone"]
    assert got[0][1] == ["Show real-world moments: a customer pulling in.", "Use bright visuals."]
